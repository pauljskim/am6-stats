"""
1회용: 구글 시트에서 내려받은 xlsx의 RawSheet를 seasons/<시즌>/history.tsv로 고정 저장하고,
Summary_sheet의 회원 명단을 config.yaml 초안으로 뽑는다.

    python3 export_history.py "data/AM6 25-26 Season Data_v2_R29.xlsx" 25-26

이후 새 경기는 seasons/<시즌>/games/ 폴더에만 추가하면 된다 (history.tsv에 없는 GameID만 games/*.yaml에서 읽음).
"""
import sys
import datetime
import openpyxl

HEADER = ['GameID', 'Date', 'Team', 'Player', 'MatchType', 'Quater',
          'Position', 'Win', 'Goal', 'Assist', 'CleanSheet',
          'Referee', 'Assistant Referee', 'Attendance', 'Venue', 'Etc.']


def clean(v, col):
    if v is None or str(v) in ('None', 'null'):
        return ''
    if col == 'MatchType' and isinstance(v, datetime.time):
        return '%d:%d' % (v.hour, v.minute)          # 엑셀이 11:11을 시각으로 바꿔버린 것 복원
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return ' '.join(str(v).split())                 # 셀 안 줄바꿈·탭 제거


def main():
    src = sys.argv[1]
    wb = openpyxl.load_workbook(src, data_only=True)

    rows = []
    for r in wb['RawSheet'].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        rows.append([clean(v, h) for v, h in zip(r[:16], HEADER)])
    out = 'seasons/%s/history.tsv' % sys.argv[2]
    with open(out, 'w', encoding='utf-8') as f:
        f.write('\t'.join(HEADER) + '\n')
        for row in rows:
            f.write('\t'.join(row) + '\n')
    print('%s: %d행' % (out, len(rows)))

    lines = []
    for r in wb['Summary_sheet'].iter_rows(min_row=2, max_row=60, values_only=True):
        if not r[1] or r[1] == '항목':
            break
        no = clean(r[0], '')
        lines.append('  - {name: %s, number: "%s", status: %s}' % (r[1], no, r[2]))
    print('\n'.join(lines))


if __name__ == '__main__':
    main()
